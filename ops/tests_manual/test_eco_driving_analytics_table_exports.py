#!/usr/bin/env python3
"""Every table in the Analizy panel downloads what it shows.

Eco Driving had exactly one export — the contributing-trip table. The ranking,
the driver-detail composition and progression tables, and the landing page's
period index had none, so the panel's main result table was the one thing a
person could read and not take away.

WHAT THIS SUITE PROTECTS is not "a route exists". It is the property that makes
an export worth having:

* **the file equals the view.** Every column the screen renders is in the file,
  in the same order; the file is the whole filtered result rather than the
  visible page; and for the ranking, whose cells are also copyable, the exported
  value of a cell is byte-identical to the value that cell copies. One
  enumeration renders all three, so they cannot disagree.
* **the scope is the page's, not a guess.** The ranking and the detail tables
  exist in two period modes, and the basis mode must route to the basis service
  method. Routing it at the period-key method would widen a week selection to a
  whole month and look correct — the defect `UI-20260827-05` already cost this
  module once.
* **a table that is not on screen is not in a file.** A non-qualified period
  renders an insufficient-distance state instead of a composition table, and
  the export of that table must refuse rather than produce the metrics the
  state exists to withhold.

Run:

    cd /opt/log-platform
    env PYTHONDONTWRITEBYTECODE=1 .venv/bin/python \\
        ops/tests_manual/test_eco_driving_analytics_table_exports.py
"""
from __future__ import annotations

import csv
import io
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
for path in (ROOT, Path(__file__).resolve().parent):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from api.eco_driving_explorer import detail_view_models as D  # noqa: E402
from api.eco_driving_explorer import eco_view as V  # noqa: E402
from api.eco_driving_explorer import ranking_view_models as R  # noqa: E402
from api.eco_driving_explorer import trip_export as X  # noqa: E402
from api.eco_driving_explorer.pages import EcoDrivingPages, _ExportFile  # noqa: E402
from api.eco_driving_explorer.service import ApiResult  # noqa: E402

CC, FAMILY, AID, MONTH, PERIOD = "ALPHA00001", "driver", "27935", "2026-08", "M:2026-08-01"
USER = {"user_id": "u-1"}

CHECKS = 0
FAILURES: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    global CHECKS
    CHECKS += 1
    if condition:
        print(f"PASS: {label}")
    else:
        FAILURES.append(label)
        print(f"FAIL: {label}\n      {detail}")


# --- fixtures ----------------------------------------------------------------


def _entry(index: int, *, qualified: bool = True) -> dict:
    return {
        "ranking_position": index + 1,
        "assigned_id": f"TAG-{index:03d}",
        "current_chart": {"current_driver_name": f"Kierowca {index}"},
        "eco_driving_score_total": str(90 - index),
        "ecodriving_rating_type": "GOOD",
        "total_distance_meters": 36000 + index * 1000,
        "trips_count": 10 + index,
        "qualification_status": "QUALIFIED" if qualified else "NOT_QUALIFIED",
        "total_distance_meters_qualifying": 36000,
        "event_counts": {m: (i + index) for i, m in enumerate(V.RANKING_METRIC_ORDER)},
        "metric_rates_per_100km": {m: "1.50" for m in V.RANKING_METRIC_ORDER},
        "metric_points": {m: 4 for m in V.RANKING_METRIC_ORDER},
        "metric_points_lost": {m: -(i + 1) for i, m in enumerate(V.RANKING_METRIC_ORDER)},
        "ranking_group": "INCLUDED",
        "capabilities": {"can_view_trip_details": True},
    }


#: More than one export page, so the paging loop is genuinely exercised.
PERIOD_ENTRIES = [_entry(i) for i in range(7)]
BASIS_ENTRIES = [_entry(i) for i in range(3)]

PROGRESSION = [
    {"period_label": "W1", "eco_driving_score_total": "80", "qualification_status": "QUALIFIED",
     "total_distance_meters": 10000, "trips_count": 3, "is_current": False},
    {"period_label": "W2", "eco_driving_score_total": "86", "qualification_status": "QUALIFIED",
     "total_distance_meters": 22000, "trips_count": 7, "is_current": False},
    {"period_label": "W3", "eco_driving_score_total": None, "qualification_status": "NOT_QUALIFIED",
     "total_distance_meters": 3000, "trips_count": 1, "is_current": True},
]

TRIPS = [
    {"provider_trip_id": f"T-{i}", "vehicle_registration": f"WX{i:04d}",
     "trip_start_ts": None, "trip_end_ts": None, "trip_distance_meters": 4000 + i,
     "assignment_source": "GPS", "source_trip_present": True,
     "total_scoring_events": i, "event_counts": {m: i for m in V.RANKING_METRIC_ORDER}}
    for i in range(4)
]

PERIODS = [
    {"period_key": f"W:2026-08-{d:02d}", "period_label": f"Tydzień {d}",
     "period_start_date": f"2026-08-{d:02d}", "period_end_date_exclusive": f"2026-08-{d + 7:02d}",
     "month_start_date": "2026-08-01", "not_ranked_count": d,
     "entry_counts_by_group": {"INCLUDED": 10 + d, "EXCLUDED": d, "UNKNOWN_DRIVER": 1}}
    for d in (3, 10, 17)
]


def _canonical_weeks(weeks):
    if not weeks:
        return weeks
    return ",".join(sorted(str(weeks).split(","), key=int))


class FakeService:
    """Answers each question with its own rows, because they ARE different.

    A fake that returned one row set for every method would let a basis export
    routed at the period-key method look correct — the exact defect this module
    has already shipped once.
    """

    def __init__(self, *, entry_qualified: bool = True) -> None:
        self.calls: list[str] = []
        self.entry_qualified = entry_qualified

    @staticmethod
    def _page(rows, page, limit, **meta):
        page, limit = int(page or 1), int(limit or 50)
        start = (page - 1) * limit
        window = rows[start:start + limit]
        return ApiResult(200, {"data": window, "meta": {
            "page": page, "limit": limit, "count": len(window),
            "total_count": len(rows), "has_next": start + limit < len(rows), **meta,
        }})

    # -- ranking ------------------------------------------------------------
    def list_ranking_entries(self, *, page=None, limit=None, **_kw):
        self.calls.append("period_key_ranking")
        return self._page(PERIOD_ENTRIES, page, limit)

    def get_basis_ranking(self, *, month=None, weeks=None, page=None, limit=None, **_kw):
        self.calls.append(f"basis_ranking(month={month},weeks={weeks})")
        return self._page(BASIS_ENTRIES, page, limit, selection={
            # The real service CANONICALIZES the week list; a fake that echoed
            # the submitted order would make the assertion below meaningless.
            "month": month or MONTH, "canonical_weeks_param": _canonical_weeks(weeks),
            "mode": "WEEKS" if weeks else "MONTH", "label": "W1+W2" if weeks else "cały miesiąc",
        }, basis={"population_count": 9, "counts_by_group": {"INCLUDED": 3},
                  "not_ranked_count": 0})

    # -- driver detail ------------------------------------------------------
    def get_ranking_entry(self, **_kw):
        self.calls.append("period_key_entry")
        return ApiResult(200, {"data": _entry(0, qualified=self.entry_qualified)})

    def get_basis_ranking_entry(self, *, month=None, weeks=None, **_kw):
        self.calls.append(f"basis_entry(month={month},weeks={weeks})")
        data = _entry(1, qualified=self.entry_qualified)
        data["selection"] = {"month": month or MONTH, "canonical_weeks_param": weeks,
                             "mode": "WEEKS" if weeks else "MONTH", "label": "W1+W2"}
        return ApiResult(200, {"data": data})

    def get_period_progression(self, *, period_key=None, **_kw):
        self.calls.append(f"progression({period_key})")
        return ApiResult(200, {"data": PROGRESSION})

    def get_driver_trend(self, **_kw):
        return ApiResult(200, {"data": []})

    def get_score_distribution(self, **_kw):
        return ApiResult(200, {"data": None})

    # -- landing ------------------------------------------------------------
    def list_providers(self, **_kw):
        return ApiResult(200, {"data": [{"client_code": CC, "ranking_family": FAMILY}]})

    def list_periods(self, *, period_type=None, **_kw):
        self.calls.append(f"periods({period_type})")
        return ApiResult(200, {"data": PERIODS, "meta": {"period_type": period_type or "WEEKLY"}})

    def list_basis_months(self, **_kw):
        return ApiResult(200, {"data": [{"month": MONTH}]})

    def list_contributing_trips(self, *, page=None, limit=None, **_kw):
        return self._page(TRIPS, page, limit)


def _csv_rows(export: _ExportFile) -> list[list[str]]:
    text = export.body.decode("utf-8")
    assert text.startswith("﻿"), "the BOM Polish Excel needs is missing"
    return [row for row in csv.reader(io.StringIO(text[1:]), delimiter=";") if row]


def _xlsx_rows(body: bytes) -> list[list]:
    """The first sheet as rows of openpyxl cells, so type AND format are visible."""
    from openpyxl import load_workbook

    return [list(row) for row in load_workbook(io.BytesIO(body)).active.iter_rows()]


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _pages(service=None) -> EcoDrivingPages:
    return EcoDrivingPages(service or FakeService())


def _hrefs(html: str, path: str) -> list[str]:
    return [h.replace("&amp;", "&") for h in re.findall(rf'href="([^"]*{path}[^"]*)"', html)]


# ===========================================================================
print("== 1. The ranking file IS the ranking table ==")
# ===========================================================================
service = FakeService()
grid = EcoDrivingPages._ranking_table(
    _pages(service), PERIOD_ENTRIES[:1], "INCLUDED",
    {"client_code": CC, "ranking_family": FAMILY, "period_key": PERIOD, "unit": "rate"},
    "eco_driving_score_total", "desc", "rate",
)
header_labels = [re.sub(r"<[^>]+>", "", cell).strip()
                 for cell in re.findall(r"<th[^>]*>(.*?)</th>", grid.html, re.S)]
# The last header is the action column; the sorted one carries a caret.
header_labels = [re.sub(r"[↑↓]\s*$", "", label).strip() for label in header_labels[:-1]]
cols = R.columns("rate")
check("every ranking column that renders a header is in the enumeration",
      len(header_labels) == len(cols), f"{len(header_labels)} headers vs {len(cols)} columns")

export = _pages(service).rankings_export(
    user=USER, export_format="csv", client_code=CC, ranking_family=FAMILY,
    period_key=PERIOD, ranking_group="INCLUDED", unit="rate",
)
rows = _csv_rows(export)
check("the exported header row is the enumeration's export labels",
      rows[0] == [c.export_label for c in cols], str(rows[0]))
check("the metric heading carries the unit the screen had under it",
      any("/ 100 km" in label for label in rows[0]), str(rows[0]))
check("the distance heading states its unit once: `Dystans (km)`, not `Dystans (km) (km)`",
      "Dystans (km)" in rows[0] and not any("(km) (km)" in label for label in rows[0]),
      str(rows[0]))
check("...while the screen's distance heading is unchanged",
      "Dystans (km)" in header_labels, str(header_labels))

# ===========================================================================
print()
print("== 2. A cell's exported value is the value that cell copies ==")
# ===========================================================================
entry = PERIOD_ENTRIES[0]
one_row = EcoDrivingPages._ranking_table(
    _pages(service), [entry], "INCLUDED", {"client_code": CC}, None, None, "rate").html
copied = re.findall(r'data-eco-column="([^"]+)"[^>]*data-eco-copy="([^"]*)"', one_row)
copied_by_column = dict(copied)
mismatch = [
    (col.key, copied_by_column.get(col.key), X.text_cell(col.value(entry)))
    for col in R.columns("rate")
    if col.key in copied_by_column
    and copied_by_column[col.key] != X.text_cell(col.value(entry)).replace("&", "&amp;")
]
check("clipboard and file agree on every cell of the row",
      not mismatch and len(copied_by_column) >= len(R.columns("rate")) - 1, str(mismatch))

# ===========================================================================
print()
print("== 3. The file is the whole view, and the unit is part of the answer ==")
# ===========================================================================
saved_page_size = X.PAGE_SIZE
X.PAGE_SIZE = 2  # force the paging loop
try:
    paged = _csv_rows(_pages(FakeService()).rankings_export(
        user=USER, export_format="csv", client_code=CC, ranking_family=FAMILY,
        period_key=PERIOD, unit="rate"))
finally:
    X.PAGE_SIZE = saved_page_size
check("a paged export still contains every row of the filtered view",
      len(paged) == 1 + len(PERIOD_ENTRIES), f"{len(paged)} rows")

sum_rows = _csv_rows(_pages(FakeService()).rankings_export(
    user=USER, export_format="csv", client_code=CC, ranking_family=FAMILY,
    period_key=PERIOD, unit="sum"))
rate_rows = _csv_rows(_pages(FakeService()).rankings_export(
    user=USER, export_format="csv", client_code=CC, ranking_family=FAMILY,
    period_key=PERIOD, unit="rate"))
metric_index = next(i for i, label in enumerate(sum_rows[0])
                    if "suma" in label.lower() or "Σ" in label)
check("the two units export two different metric readings, each labelled",
      sum_rows[0] != rate_rows[0] and sum_rows[1][metric_index] != rate_rows[1][metric_index],
      f"{sum_rows[0][metric_index]!r} vs {rate_rows[0][metric_index]!r}")

# ===========================================================================
print()
print("== 4. The scope is the page's, in both period modes ==")
# ===========================================================================
service = FakeService()
basis = _csv_rows(_pages(service).rankings_export(
    user=USER, export_format="csv", client_code=CC, ranking_family=FAMILY,
    month=MONTH, weeks="1,2", unit="rate"))
check("a month/weeks export is answered by the BASIS method, not the period one",
      any(c.startswith("basis_ranking(month=2026-08,weeks=1,2)") for c in service.calls)
      and "period_key_ranking" not in service.calls, str(service.calls))
check("...and contains that basis's rows, not the month's",
      len(basis) == 1 + len(BASIS_ENTRIES), f"{len(basis)} rows")

refused = _pages().rankings_export(user=USER, export_format="csv", client_code=CC,
                                   ranking_family=FAMILY, unit="rate")
check("with neither a period nor a month the export refuses rather than inventing a scope",
      not isinstance(refused, _ExportFile) and refused.status_code == 422,
      getattr(refused, "status_code", refused))

# ===========================================================================
print()
print("== 5. The ranking page offers the download, at the canonical scope ==")
# ===========================================================================
html = _pages(FakeService()).rankings(
    user=USER, client_code=CC, ranking_family=FAMILY, period_key=PERIOD,
    ranking_group="INCLUDED", unit="rate").body_html
links = _hrefs(html, "rankings/export")
check("the period-key ranking offers XLSX and CSV", len(links) == 2, str(links))
check("...naming the period and never the pagination",
      links and "period_key=" in links[0] and "page=" not in links[0], str(links[:1]))
check("...and the caption states the row count",
      f"({len(PERIOD_ENTRIES)})" in html, "caption count missing")

basis_html = _pages(FakeService()).rankings(
    user=USER, client_code=CC, ranking_family=FAMILY, month=MONTH, weeks="2,1",
    ranking_group="INCLUDED", unit="rate").body_html
basis_links = _hrefs(basis_html, "rankings/export")
check("the basis ranking's link carries the month", basis_links and "month=2026-08" in basis_links[0],
      str(basis_links[:1]))
check("...the CANONICAL weeks, not the submitted order",
      basis_links and "weeks=1%2C2" in basis_links[0] and "weeks=2%2C1" not in basis_links[0],
      str(basis_links[:1]))
check("...and no period key, which would widen it to the month",
      basis_links and "period_key=" not in basis_links[0], str(basis_links[:1]))

# ===========================================================================
print()
print("== 6. The driver-detail tables export what their panels show ==")
# ===========================================================================
service = FakeService()
comp = _csv_rows(_pages(service).ranking_entry_export(
    user=USER, export_format="csv", table="composition", client_code=CC,
    ranking_family=FAMILY, period_key=PERIOD, assigned_id=AID))
check("the composition file has one row per metric",
      len(comp) == 1 + len(D.composition_rows(_entry(0))), f"{len(comp)} rows")
check("...in the screen's order — worst loss first",
      [row[0] for row in comp[1:]] == [r["metric_label"] for r in D.composition_rows(_entry(0))],
      str([row[0] for row in comp[1:]][:3]))

prog = _csv_rows(_pages(FakeService()).ranking_entry_export(
    user=USER, export_format="csv", table="progression", client_code=CC,
    ranking_family=FAMILY, period_key=PERIOD, assigned_id=AID))
check("the progression file has one row per snapshot", len(prog) == 1 + len(PROGRESSION), f"{len(prog)} rows")
check("...and its distance heading states the unit once",
      "Dystans (km)" in prog[0] and not any("(km) (km)" in label for label in prog[0]),
      str(prog[0]))
delta_index = prog[0].index("Δ")
check("...and the delta is the running one the table computes, blank where the table dashes",
      prog[1][delta_index] == "" and prog[2][delta_index] not in ("", "—"),
      f"{prog[1][delta_index]!r} then {prog[2][delta_index]!r}")

service = FakeService()
_pages(service).ranking_entry_export(
    user=USER, export_format="csv", table="progression", client_code=CC,
    ranking_family=FAMILY, month=MONTH, weeks="1,2", assigned_id=AID)
check("a basis progression export reads the basis entry and the month's history",
      any(c.startswith("basis_entry(") for c in service.calls)
      and "period_key_entry" not in service.calls, str(service.calls))

unknown = _pages().ranking_entry_export(
    user=USER, export_format="csv", table="trend", client_code=CC,
    ranking_family=FAMILY, period_key=PERIOD, assigned_id=AID)
check("an unnamed table is refused rather than defaulted",
      not isinstance(unknown, _ExportFile) and unknown.status_code == 400,
      getattr(unknown, "status_code", unknown))

blocked = _pages(FakeService(entry_qualified=False)).ranking_entry_export(
    user=USER, export_format="csv", table="composition", client_code=CC,
    ranking_family=FAMILY, period_key=PERIOD, assigned_id=AID)
check("a non-qualified period has no composition table and exports none",
      not isinstance(blocked, _ExportFile) and blocked.status_code == 409,
      getattr(blocked, "status_code", blocked))

detail_html = _pages(FakeService()).ranking_entry(
    user=USER, client_code=CC, ranking_family=FAMILY, period_key=PERIOD,
    assigned_id=AID, unit="rate").body_html
detail_links = _hrefs(detail_html, "ranking-entry/export")
trips_export_links = _hrefs(detail_html, "trips/export")
check("the trip-evidence panel offers the download too, through the full trip endpoint",
      len(trips_export_links) == 2
      and all("assigned_id=" in h and "period_key=" in h for h in trips_export_links),
      str(trips_export_links))

check("both detail panels offer both formats",
      len(detail_links) == 4
      and sum("table=composition" in h for h in detail_links) == 2
      and sum("table=progression" in h for h in detail_links) == 2,
      str(detail_links))

# ===========================================================================
print()
print("== 7. The landing page's period index exports too ==")
# ===========================================================================
periods = _csv_rows(_pages(FakeService()).periods_export(
    user=USER, export_format="csv", client_code=CC, ranking_family=FAMILY))
check("the period file has one row per period", len(periods) == 1 + len(PERIODS), f"{len(periods)} rows")
check("...and carries the group counts the screen shows, as separate columns",
      periods[1][3:6] == ["13", "3", "1"], str(periods[1]))

landing_html = _pages(FakeService()).landing(
    user=USER, client_code=CC, ranking_family=FAMILY).body_html
landing_links = _hrefs(landing_html, "periods/export")
check("the landing page offers both formats above its table", len(landing_links) == 2, str(landing_links))

# ===========================================================================
print()
print("== 8. Format contract ==")
# ===========================================================================
xlsx = _pages(FakeService()).rankings_export(
    user=USER, export_format="xlsx", client_code=CC, ranking_family=FAMILY,
    period_key=PERIOD, unit="rate")
check("XLSX is a real workbook", isinstance(xlsx, _ExportFile) and xlsx.body[:2] == b"PK",
      str(xlsx.body[:8]))
check("...declared as one", xlsx.content_type.startswith("application/vnd.openxmlformats"),
      xlsx.content_type)
check("the scope rides into the filename parts, naming the period",
      any(name == "period_key" for name, _ in xlsx.scope), str(xlsx.scope))
sheet = _xlsx_rows(xlsx.body)
check("the XLSX header row is the CSV header row",
      [c.value for c in sheet[0]] == rate_rows[0], str([c.value for c in sheet[0]]))

# ===========================================================================
print()
print("== 9. XLSX metric cells are numbers; the clipboard and CSV text did not move ==")
# ===========================================================================
# The text a metric cell copies was defined by `metric_copy_value` before the
# column's value became a number. `text_cell(metric_number(...))` must still
# produce it, byte for byte — including half-way rates, where a float rounding
# rule would disagree with the Decimal one.
RATES = ["1.50", "0", "0.00", "1.125", "1.135", "0.005", "2.999", "123.455", "7", "10.1", None, ""]
rate_drift = [
    (rate, X.text_cell(V.metric_number(unit="rate", event_count=None, rate=rate)),
     V.metric_copy_value(unit="rate", event_count=None, rate=rate))
    for rate in RATES
]
rate_drift = [d for d in rate_drift if d[1] != d[2]]
check("rate: clipboard text of the number == the text the cell copied before", not rate_drift,
      str(rate_drift))
sum_drift = [
    (count, X.text_cell(V.metric_number(unit="sum", event_count=count, rate=None)),
     V.metric_copy_value(unit="sum", event_count=count, rate=None))
    for count in (0, 1, 17, 12345, None)
]
sum_drift = [d for d in sum_drift if d[1] != d[2]]
check("Σ: clipboard text of the number == the text the cell copied before", not sum_drift,
      str(sum_drift))

below = _entry(1, qualified=False)
mixed = [dict(_entry(0), metric_rates_per_100km={
    m: RATES[i % 10] for i, m in enumerate(V.RANKING_METRIC_ORDER)}), below]
metric_keys = list(V.RANKING_METRIC_ORDER)
for unit, number_format in (("rate", "0.00"), ("sum", "General")):
    cols = R.columns(unit)
    metric_idx = [i for i, c in enumerate(cols) if c.key in metric_keys]
    score_idx = next(i for i, c in enumerate(cols) if c.key == "eco_driving_score_total")
    book = _xlsx_rows(X.to_xlsx(cols, mixed, sheet_title="Ranking"))
    qualified_cells = [book[1][i] for i in metric_idx]
    check(f"{unit}: every metric cell of a qualified row is a number, not text",
          len(qualified_cells) == len(metric_keys)
          and all(_is_number(c.value) for c in qualified_cells),
          str([(type(c.value).__name__, c.value) for c in qualified_cells]))
    check(f"{unit}: ...displayed as `{number_format}`",
          all(c.number_format == number_format for c in qualified_cells),
          str([c.number_format for c in qualified_cells]))
    expected = [V.metric_number(unit=unit, event_count=mixed[0]["event_counts"][m],
                                rate=mixed[0]["metric_rates_per_100km"][m]) for m in metric_keys]
    check(f"{unit}: ...holding the value the screen shows",
          [c.value for c in qualified_cells] == expected,
          f"{[c.value for c in qualified_cells]} vs {expected}")
    if unit == "sum":
        check("Σ: the cells are integers, the stored event counters",
              [c.value for c in qualified_cells] == [mixed[0]["event_counts"][m] for m in metric_keys]
              and all(isinstance(c.value, int) for c in qualified_cells),
              str([c.value for c in qualified_cells]))
    check(f"{unit}: a below-threshold row exports empty metric and score cells",
          all(book[2][i].value is None for i in metric_idx + [score_idx]),
          str([book[2][i].value for i in metric_idx + [score_idx]]))
    text = [row for row in csv.reader(
        io.StringIO(X.to_csv(cols, mixed).decode("utf-8")[1:]), delimiter=";") if row]
    copied = [V.metric_copy_value(unit=unit, event_count=mixed[0]["event_counts"][m],
                                  rate=mixed[0]["metric_rates_per_100km"][m]) for m in metric_keys]
    check(f"{unit}: the CSV still carries the copy text, and blanks below the threshold",
          [text[1][i] for i in metric_idx] == copied
          and all(text[2][i] == "" for i in metric_idx + [score_idx]),
          f"{[text[1][i] for i in metric_idx]} vs {copied}")

print()
if FAILURES:
    print(f"FAILED {len(FAILURES)}/{CHECKS}: " + "; ".join(FAILURES))
    raise SystemExit(1)
print(f"OK - Eco Driving Analizy table export checks passed ({CHECKS} checks)")
