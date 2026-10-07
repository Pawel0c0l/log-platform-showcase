"""
Deprecated direct ALPHA00001 Alpha GPS workbook import.

The server never executes VBA/macros. It reads the saved XLSM workbook with
openpyxl in read-only/data-only mode, validates the LOG worksheet, then
transactionally replaces telematics_reports."Alpha_GPS_Baza_LOG" for a new
source sha256.

Canonical operations now use Workflow B email ingestion for report type
Alpha_GPS_Baza_LOG. Keep this module only for emergency/manual compatibility
with the previously implemented source_path flow.
"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
import time
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from jobs.api.telematics.secret_resolver import resolve_secret


JOB_SOURCE = "jobs.alpha.import_gps_baza_log_xlsm"
DEPRECATED = True
WORKFLOW_NAME = "workflow_a"
STAGE_NAME = "alpha_gps_baza_log_import"
DEFAULT_CLIENT_CODE = "ALPHA00001"
EXPECTED_CLIENT_DB_NAME = "alpha_main"
DEFAULT_SHEET_NAME = "LOG"
DESTINATION_SCHEMA = "telematics_reports"
DESTINATION_TABLE = "Alpha_GPS_Baza_LOG"
IMPORT_RUNS_TABLE = "alpha_gps_baza_log_import_runs"
REQUIRED_HEADERS = ("ID", "Nr rejestracyjny", "Data przydziału", "Nazwa Pliku csv")


@dataclass(frozen=True)
class ClientDbConfig:
    client_id: str
    client_code: str
    client_db_host: str
    client_db_port: int
    client_db_name: str
    client_db_user: str
    client_db_password_secret_ref: str


@dataclass(frozen=True)
class SourceFileInfo:
    path: Path
    size_bytes: int
    modified_ns: int
    sha256: str


@dataclass(frozen=True)
class ParsedRow:
    source_id: Optional[str]
    registration: Optional[str]
    assignment_date: Optional[date]
    csv_filename: Optional[str]
    source_row_number: int
    raw_row_json: dict[str, Any]


@dataclass(frozen=True)
class ParsedWorkbook:
    sheet_name: str
    header_row_number: int
    rows: list[ParsedRow]
    empty_rows_skipped: int


def _require_dependency(module: str, feature: str):
    try:
        return __import__(module)
    except ImportError as exc:
        raise RuntimeError(f"Missing dependency for {feature}: {module}") from exc


def _bool_param(value: Any, default: bool = False) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return bool(value)
    text = str(value).strip().lower()
    if text in {"1", "true", "t", "yes", "y", "on"}:
        return True
    if text in {"0", "false", "f", "no", "n", "off"}:
        return False
    return default


def _platform_pg_conn():
    psycopg = _require_dependency("psycopg", "Postgres connection")
    from psycopg.rows import dict_row

    dsn = (
        f"host={os.getenv('POSTGRES_HOST', '127.0.0.1')} "
        f"port={os.getenv('POSTGRES_PORT', '5432')} "
        f"dbname={os.getenv('POSTGRES_DB', 'logdb')} "
        f"user={os.getenv('POSTGRES_USER', 'loguser')} "
        f"password={os.getenv('POSTGRES_PASSWORD', '')}"
    )
    return psycopg.connect(dsn, row_factory=dict_row)


def _client_business_pg_conn(config: ClientDbConfig):
    psycopg = _require_dependency("psycopg", "Postgres connection")
    from psycopg.rows import dict_row

    return psycopg.connect(
        host=config.client_db_host,
        port=config.client_db_port,
        dbname=config.client_db_name,
        user=config.client_db_user,
        password=resolve_secret(config.client_db_password_secret_ref),
        row_factory=dict_row,
        autocommit=False,
    )


def _load_client_db_config(*, client_code: str) -> ClientDbConfig:
    conn = _platform_pg_conn()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT
                    client_id::text AS client_id,
                    client_code,
                    client_db_host,
                    client_db_port,
                    client_db_name,
                    client_db_user,
                    client_db_password_secret_ref
                FROM workflow_a_control.client_account
                WHERE enabled = true
                  AND client_code = %s
                LIMIT 1
                """,
                (client_code,),
            )
            row = cur.fetchone()
            if not row:
                raise RuntimeError(f"Enabled client_account not found for client_code={client_code}")
            client_db_name = str(row["client_db_name"])
            if client_code == DEFAULT_CLIENT_CODE and client_db_name != EXPECTED_CLIENT_DB_NAME:
                raise RuntimeError(
                    f"ALPHA00001 must resolve to client_db_name={EXPECTED_CLIENT_DB_NAME!r}; "
                    f"got {client_db_name!r}"
                )
            return ClientDbConfig(
                client_id=str(row["client_id"]),
                client_code=str(row["client_code"]),
                client_db_host=str(row["client_db_host"]),
                client_db_port=int(row["client_db_port"]),
                client_db_name=client_db_name,
                client_db_user=str(row["client_db_user"]),
                client_db_password_secret_ref=str(row["client_db_password_secret_ref"]),
            )
    finally:
        conn.close()


def _stable_source_file(path: Path, *, sleep_s: float = 1.0) -> SourceFileInfo:
    if not path.exists():
        raise FileNotFoundError(f"source_path does not exist: {path}")
    if not path.is_file():
        raise ValueError(f"source_path is not a file: {path}")

    first = path.stat()
    if sleep_s > 0:
        time.sleep(sleep_s)
    second = path.stat()
    if first.st_size != second.st_size or first.st_mtime_ns != second.st_mtime_ns:
        raise RuntimeError(
            "source_path changed while checking file stability; refusing to read a partial copy"
        )

    sha = _sha256_file(path)
    return SourceFileInfo(
        path=path,
        size_bytes=int(second.st_size),
        modified_ns=int(second.st_mtime_ns),
        sha256=sha,
    )


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _trim_string(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_registration(value: Any) -> Optional[str]:
    text = _trim_string(value)
    if text is None:
        return None
    return " ".join(text.split()).upper()


def _json_safe(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _row_value(row: tuple[Any, ...], pos: int) -> Any:
    if pos < 0 or pos >= len(row):
        return None
    return row[pos]


def _header_values(row: tuple[Any, ...]) -> list[str]:
    return [str(value).strip() for value in row if value is not None and str(value).strip()]


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = _header_values(row)
        if values == list(REQUIRED_HEADERS):
            positions: dict[str, int] = {}
            for idx, value in enumerate(row):
                text = str(value).strip() if value is not None else ""
                if text in REQUIRED_HEADERS:
                    positions[text] = idx
            missing = [h for h in REQUIRED_HEADERS if h not in positions]
            if missing:
                raise ValueError(f"Required columns missing from header row {row_idx}: {missing}")
            return row_idx, positions
    raise ValueError(
        "Required header row not found; expected non-empty cells exactly: "
        + ", ".join(REQUIRED_HEADERS)
    )


def _parse_assignment_date(value: Any, *, row_number: int) -> Optional[date]:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel

            return from_excel(value).date()
        except Exception as exc:
            raise ValueError(f"Invalid Data przydziału at source row {row_number}: {value!r}") from exc

    text = str(value).strip()
    formats = (
        "%Y-%m-%d",
        "%d.%m.%Y",
        "%d/%m/%Y",
        "%Y/%m/%d",
        "%d-%m-%Y",
        "%Y-%m-%d %H:%M:%S",
        "%d.%m.%Y %H:%M:%S",
    )
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date()
    except ValueError as exc:
        raise ValueError(f"Invalid Data przydziału at source row {row_number}: {text!r}") from exc


def _read_workbook(path: Path, *, sheet_name: str) -> ParsedWorkbook:
    openpyxl = _require_dependency("openpyxl", "XLSM workbook import")

    workbook = openpyxl.load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
        keep_vba=False,
    )
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet_name}")
        ws = workbook[sheet_name]
        header_row_number, header_positions = _find_header_row(ws)

        rows: list[ParsedRow] = []
        empty_rows_skipped = 0
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row_idx <= header_row_number:
                continue
            raw_values = {
                header: _json_safe(row[pos] if pos < len(row) else None)
                for header, pos in header_positions.items()
            }
            if all(value is None or str(value).strip() == "" for value in raw_values.values()):
                empty_rows_skipped += 1
                continue

            assignment_date = _parse_assignment_date(
                row[header_positions["Data przydziału"]]
                if header_positions["Data przydziału"] < len(row)
                else None,
                row_number=row_idx,
            )
            rows.append(
                ParsedRow(
                    source_id=_trim_string(_row_value(row, header_positions["ID"])),
                    registration=_normalize_registration(
                        _row_value(row, header_positions["Nr rejestracyjny"])
                    ),
                    assignment_date=assignment_date,
                    csv_filename=_trim_string(_row_value(row, header_positions["Nazwa Pliku csv"])),
                    source_row_number=row_idx,
                    raw_row_json={
                        "source_row_number": row_idx,
                        "values": raw_values,
                    },
                )
            )
        return ParsedWorkbook(
            sheet_name=sheet_name,
            header_row_number=header_row_number,
            rows=rows,
            empty_rows_skipped=empty_rows_skipped,
        )
    finally:
        workbook.close()


def _pg_sql():
    _require_dependency("psycopg", "Postgres SQL composition")
    from psycopg import sql as pgsql

    return pgsql


def _pg_jsonb(value: Any):
    _require_dependency("psycopg", "Postgres JSONB")
    from psycopg.types.json import Jsonb

    return Jsonb(value)


def _successful_sha_exists(conn, source_sha256: str) -> Optional[dict[str, Any]]:
    sql = _pg_sql()
    query = sql.SQL(
        "SELECT import_run_id::text, rows_loaded, finished_at "
        "FROM {schema}.{history} "
        "WHERE source_sha256 = %s AND status = 'SUCCESS' "
        "ORDER BY finished_at DESC NULLS LAST "
        "LIMIT 1"
    ).format(
        schema=sql.Identifier(DESTINATION_SCHEMA),
        history=sql.Identifier(IMPORT_RUNS_TABLE),
    )
    with conn.cursor() as cur:
        cur.execute(query, (source_sha256,))
        row = cur.fetchone()
        return dict(row) if row else None


def _insert_import_run(
    conn,
    *,
    import_run_id: str,
    source: SourceFileInfo,
    metadata: dict[str, Any],
) -> None:
    sql = _pg_sql()
    query = sql.SQL(
        "INSERT INTO {schema}.{history} ("
        " import_run_id, source_path, source_filename, source_sha256,"
        " source_size_bytes, status, rows_loaded, metadata_json"
        ") VALUES (%s, %s, %s, %s, %s, 'RUNNING', 0, %s)"
    ).format(
        schema=sql.Identifier(DESTINATION_SCHEMA),
        history=sql.Identifier(IMPORT_RUNS_TABLE),
    )
    with conn.cursor() as cur:
        cur.execute(
            query,
            (
                import_run_id,
                str(source.path),
                source.path.name,
                source.sha256,
                source.size_bytes,
                _pg_jsonb(metadata),
            ),
        )
    conn.commit()


def _mark_import_run_failed(conn, *, import_run_id: str, error_message: str, metadata: dict[str, Any]) -> None:
    sql = _pg_sql()
    query = sql.SQL(
        "UPDATE {schema}.{history} "
        "SET status = 'FAILED', finished_at = now(), error_message = %s, metadata_json = %s "
        "WHERE import_run_id = %s"
    ).format(
        schema=sql.Identifier(DESTINATION_SCHEMA),
        history=sql.Identifier(IMPORT_RUNS_TABLE),
    )
    with conn.cursor() as cur:
        cur.execute(query, (error_message[:4000], _pg_jsonb(metadata), import_run_id))
    conn.commit()


def _replace_target_table(
    conn,
    *,
    import_run_id: str,
    source: SourceFileInfo,
    workbook: ParsedWorkbook,
    metadata: dict[str, Any],
) -> None:
    sql = _pg_sql()
    target = sql.SQL("{}.{}").format(
        sql.Identifier(DESTINATION_SCHEMA),
        sql.Identifier(DESTINATION_TABLE),
    )
    history = sql.SQL("{}.{}").format(
        sql.Identifier(DESTINATION_SCHEMA),
        sql.Identifier(IMPORT_RUNS_TABLE),
    )
    delete_sql = sql.SQL("DELETE FROM {}").format(target)
    insert_sql = sql.SQL(
        "INSERT INTO {} ("
        " source_id, registration, assignment_date, csv_filename,"
        " source_sha256, source_row_number, raw_row_json"
        ") VALUES (%s, %s, %s, %s, %s, %s, %s)"
    ).format(target)
    success_sql = sql.SQL(
        "UPDATE {} "
        "SET status = 'SUCCESS', finished_at = now(), rows_loaded = %s,"
        "    error_message = NULL, metadata_json = %s "
        "WHERE import_run_id = %s"
    ).format(history)
    supersede_sql = sql.SQL(
        "UPDATE {} "
        "SET status = 'SUPERSEDED', metadata_json = COALESCE(metadata_json, '{}'::jsonb) || %s "
        "WHERE source_sha256 = %s AND status = 'SUCCESS' AND import_run_id <> %s"
    ).format(history)

    with conn.cursor() as cur:
        cur.execute(delete_sql)
        cur.executemany(
            insert_sql,
            [
                (
                    row.source_id,
                    row.registration,
                    row.assignment_date,
                    row.csv_filename,
                    source.sha256,
                    row.source_row_number,
                    _pg_jsonb(row.raw_row_json),
                )
                for row in workbook.rows
            ],
        )
        cur.execute(
            supersede_sql,
            (
                _pg_jsonb({"superseded_by_import_run_id": import_run_id}),
                source.sha256,
                import_run_id,
            ),
        )
        cur.execute(success_sql, (len(workbook.rows), _pg_jsonb(metadata), import_run_id))
    conn.commit()


def _upload_artifacts(
    client,
    *,
    run_id: str,
    source: SourceFileInfo,
    summary: dict[str, Any],
    upload_source_workbook: bool,
) -> None:
    with tempfile.TemporaryDirectory(prefix="alpha-gps-baza-log-import-") as tmpdir:
        tmp = Path(tmpdir)
        summary_path = tmp / "alpha_gps_baza_log_import_summary.json"
        summary_path.write_text(json.dumps(summary, indent=2, sort_keys=True), encoding="utf-8")
        client.upload_artifact(
            str(summary_path),
            kind="JSON",
            run_id=run_id,
            workflow_name=WORKFLOW_NAME,
            stage_name=STAGE_NAME,
            artifact_role="import_summary",
            client_code=summary.get("client_code"),
            original_filename=source.path.name,
            display_filename=summary_path.name,
            metadata=summary,
        )
        if upload_source_workbook:
            workbook_copy = tmp / source.path.name
            shutil.copy2(source.path, workbook_copy)
            client.upload_artifact(
                str(workbook_copy),
                kind="REPORT_RAW",
                run_id=run_id,
                workflow_name=WORKFLOW_NAME,
                stage_name=STAGE_NAME,
                artifact_role="source_xlsm",
                client_code=summary.get("client_code"),
                original_filename=source.path.name,
                display_filename=source.path.name,
                metadata={
                    "source_path": str(source.path),
                    "source_sha256": source.sha256,
                    "source_size_bytes": source.size_bytes,
                },
            )


def _summary_base(
    *,
    cfg: ClientDbConfig,
    source: SourceFileInfo,
    sheet_name: str,
    dry_run: bool,
    force: bool,
    trigger: Any,
) -> dict[str, Any]:
    return {
        "job_source": JOB_SOURCE,
        "client_id": cfg.client_id,
        "client_code": cfg.client_code,
        "client_db_name": cfg.client_db_name,
        "destination_table": f'{DESTINATION_SCHEMA}."{DESTINATION_TABLE}"',
        "source_path": str(source.path),
        "source_filename": source.path.name,
        "source_sha256": source.sha256,
        "source_size_bytes": source.size_bytes,
        "sheet_name": sheet_name,
        "dry_run": dry_run,
        "force": force,
        "trigger": trigger,
    }


def run(client, run_id: str, params: dict):
    if not isinstance(params, dict):
        raise ValueError("params must be a dict")

    source_path_raw = params.get("source_path")
    if not source_path_raw:
        raise ValueError("Missing required param: source_path")
    sheet_name = str(params.get("sheet_name") or DEFAULT_SHEET_NAME).strip() or DEFAULT_SHEET_NAME
    client_code = str(params.get("client_code") or DEFAULT_CLIENT_CODE).strip()
    dry_run = _bool_param(params.get("dry_run"), False)
    force = _bool_param(params.get("force"), False)
    trigger = params.get("trigger", "MANUAL")
    stability_sleep_s = float(params.get("file_stability_sleep_s", 1.0))

    client.log(
        "WARNING",
        "SCRIPT",
        JOB_SOURCE,
        "Deprecated direct Alpha GPS XLSM import used; canonical path is Workflow B email ingestion",
        run_id=run_id,
        context={
            "deprecated": DEPRECATED,
            "canonical_report_type": "Alpha_GPS_Baza_LOG",
            "canonical_stage3_target": 'telematics_reports."Alpha_GPS_Baza_LOG"',
        },
    )

    source = _stable_source_file(Path(str(source_path_raw)), sleep_s=stability_sleep_s)
    cfg = _load_client_db_config(client_code=client_code)

    client.log(
        "INFO",
        "SCRIPT",
        JOB_SOURCE,
        "Alpha GPS XLSM import starting",
        run_id=run_id,
        context={
            "client_code": cfg.client_code,
            "client_db_name": cfg.client_db_name,
            "source_path": str(source.path),
            "source_sha256": source.sha256,
            "source_size_bytes": source.size_bytes,
            "sheet_name": sheet_name,
            "dry_run": dry_run,
            "force": force,
            "trigger": trigger,
        },
    )

    conn = _client_business_pg_conn(cfg)
    import_run_id = str(uuid.uuid4())
    import_run_started = False
    summary: dict[str, Any] = _summary_base(
        cfg=cfg,
        source=source,
        sheet_name=sheet_name,
        dry_run=dry_run,
        force=force,
        trigger=trigger,
    )
    try:
        duplicate = _successful_sha_exists(conn, source.sha256)
        if duplicate and not force:
            summary.update(
                {
                    "status": "SKIPPED_DUPLICATE_SHA256",
                    "skipped_duplicate_sha": True,
                    "previous_import_run_id": duplicate["import_run_id"],
                    "previous_rows_loaded": duplicate["rows_loaded"],
                    "rows_loaded": 0,
                }
            )
            client.log(
                "INFO",
                "SCRIPT",
                JOB_SOURCE,
                "Alpha GPS XLSM import skipped; sha256 already imported successfully",
                run_id=run_id,
                context=summary,
            )
            try:
                _upload_artifacts(
                    client,
                    run_id=run_id,
                    source=source,
                    summary=summary,
                    upload_source_workbook=False,
                )
            except Exception as exc:
                client.log(
                    "WARNING",
                    "SCRIPT",
                    JOB_SOURCE,
                    "Failed to upload duplicate-skip summary artifact",
                    run_id=run_id,
                    context={"error": str(exc), "source_sha256": source.sha256},
                )
            return summary

        if not dry_run:
            summary.update(
                {
                    "import_run_id": import_run_id,
                    "status": "RUNNING",
                    "skipped_duplicate_sha": False,
                }
            )
            _insert_import_run(conn, import_run_id=import_run_id, source=source, metadata=summary)
            import_run_started = True
            client.log(
                "INFO",
                "SCRIPT",
                JOB_SOURCE,
                "Alpha GPS XLSM import history row created",
                run_id=run_id,
                context={
                    "import_run_id": import_run_id,
                    "client_code": cfg.client_code,
                    "source_path": str(source.path),
                    "source_sha256": source.sha256,
                    "source_size_bytes": source.size_bytes,
                    "dry_run": dry_run,
                    "force": force,
                },
            )

        try:
            workbook = _read_workbook(source.path, sheet_name=sheet_name)
        except Exception as exc:
            if import_run_started:
                failed_summary = dict(summary)
                failed_summary.update({"status": "FAILED", "error_message": str(exc)})
                _mark_import_run_failed(
                    conn,
                    import_run_id=import_run_id,
                    error_message=str(exc),
                    metadata=failed_summary,
                )
            raise

        summary.update(
            {
                "import_run_id": import_run_id,
                "status": "DRY_RUN" if dry_run else "RUNNING",
                "detected_worksheet": workbook.sheet_name,
                "header_row_number": workbook.header_row_number,
                "input_rows": len(workbook.rows),
                "empty_rows_skipped": workbook.empty_rows_skipped,
                "skipped_duplicate_sha": False,
            }
        )

        if dry_run:
            summary.update(
                {
                    "status": "DRY_RUN_SUCCESS",
                    "rows_loaded": 0,
                    "would_replace_table": True,
                    "would_load_rows": len(workbook.rows),
                }
            )
            client.log(
                "INFO",
                "SCRIPT",
                JOB_SOURCE,
                "Alpha GPS XLSM dry-run validated",
                run_id=run_id,
                context=summary,
            )
            try:
                _upload_artifacts(
                    client,
                    run_id=run_id,
                    source=source,
                    summary=summary,
                    upload_source_workbook=False,
                )
            except Exception as exc:
                client.log(
                    "WARNING",
                    "SCRIPT",
                    JOB_SOURCE,
                    "Failed to upload dry-run summary artifact",
                    run_id=run_id,
                    context={"error": str(exc), "source_sha256": source.sha256},
                )
            return summary

        try:
            summary["status"] = "SUCCESS"
            summary["rows_loaded"] = len(workbook.rows)
            _replace_target_table(
                conn,
                import_run_id=import_run_id,
                source=source,
                workbook=workbook,
                metadata=summary,
            )
        except Exception as exc:
            try:
                conn.rollback()
            except Exception:
                pass
            failed_summary = dict(summary)
            failed_summary.update({"status": "FAILED", "error_message": str(exc)})
            _mark_import_run_failed(
                conn,
                import_run_id=import_run_id,
                error_message=str(exc),
                metadata=failed_summary,
            )
            raise

        client.log(
            "INFO",
            "SCRIPT",
            JOB_SOURCE,
            "Alpha GPS XLSM import completed",
            run_id=run_id,
            context=summary,
        )
        try:
            _upload_artifacts(
                client,
                run_id=run_id,
                source=source,
                summary=summary,
                upload_source_workbook=True,
            )
        except Exception as exc:
            client.log(
                "WARNING",
                "SCRIPT",
                JOB_SOURCE,
                "Failed to upload import artifacts",
                run_id=run_id,
                context={"error": str(exc), "source_sha256": source.sha256},
            )
        return summary
    except Exception:
        try:
            conn.rollback()
        except Exception:
            pass
        raise
    finally:
        conn.close()
