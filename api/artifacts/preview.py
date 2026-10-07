from __future__ import annotations

import csv
import io
import json
import os
from typing import Any


ARTIFACT_PREVIEW_MAX_BYTES = 25 * 1024 * 1024
ARTIFACT_PREVIEW_ROWS_DEFAULT = 1000
ARTIFACT_PREVIEW_ROWS_MAX = 5000
ARTIFACT_PREVIEW_TEXT_CHARS_DEFAULT = 20000
ARTIFACT_PREVIEW_TEXT_CHARS_MAX = 100000

SUPPORTED_TABLE_EXTS = {"csv", "xls", "xlsx"}
SUPPORTED_TEXT_EXTS = {"txt", "log"}


class ArtifactPreviewError(Exception):
    def __init__(self, status_code: int, detail: str):
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


def clamp_rows_limit(value: int | None) -> int:
    if value is None:
        return ARTIFACT_PREVIEW_ROWS_DEFAULT
    return max(1, min(int(value), ARTIFACT_PREVIEW_ROWS_MAX))


def clamp_text_chars_limit(value: int | None) -> int:
    if value is None:
        return ARTIFACT_PREVIEW_TEXT_CHARS_DEFAULT
    return max(1, min(int(value), ARTIFACT_PREVIEW_TEXT_CHARS_MAX))


def infer_preview_ext(artifact: dict) -> str:
    for value in (
        artifact.get("file_ext"),
        os.path.splitext(artifact.get("display_filename") or "")[1].lstrip("."),
        os.path.splitext(artifact.get("filename") or "")[1].lstrip("."),
        os.path.splitext(artifact.get("original_filename") or "")[1].lstrip("."),
    ):
        if value:
            return str(value).strip().lower().lstrip(".")

    content_type = (artifact.get("content_type") or "").split(";")[0].strip().lower()
    if content_type in {"text/csv", "application/csv"}:
        return "csv"
    if content_type in {"text/plain", "text/x-log"}:
        return "txt"
    if content_type in {"application/json", "text/json"}:
        return "json"
    if content_type == "application/pdf":
        return "pdf"
    return ""


def file_too_large_response(artifact: dict, *, max_preview_bytes: int = ARTIFACT_PREVIEW_MAX_BYTES) -> dict:
    return {
        "artifact_id": str(artifact.get("artifact_id")),
        "preview_type": "unavailable",
        "reason": "file_too_large",
        "size_bytes": artifact.get("size_bytes"),
        "max_preview_bytes": max_preview_bytes,
    }


def _decode_text(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "cp1250", "latin-1"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replace"


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _dedupe_columns(columns: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for idx, column in enumerate(columns, start=1):
        base = _stringify(column).strip() or f"column_{idx}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        result.append(base if count == 0 else f"{base}_{count + 1}")
    return result


def _rows_from_matrix(matrix: list[list[Any]], rows_limit: int) -> tuple[list[str], list[dict[str, str]], bool]:
    if not matrix:
        return [], [], False
    width = max(len(row) for row in matrix)
    padded = [row + [""] * (width - len(row)) for row in matrix]
    columns = _dedupe_columns([_stringify(value) for value in padded[0]])
    data_rows = padded[1:]
    truncated = len(data_rows) > rows_limit
    rows = [
        {columns[idx]: _stringify(value) for idx, value in enumerate(row)}
        for row in data_rows[:rows_limit]
    ]
    return columns, rows, truncated


def csv_preview(artifact: dict, data: bytes, *, rows_limit: int) -> dict:
    text, encoding = _decode_text(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    matrix: list[list[str]] = []
    for idx, row in enumerate(reader):
        matrix.append(row)
        if idx > rows_limit:
            break
    columns, rows, truncated = _rows_from_matrix(matrix, rows_limit)
    return {
        "artifact_id": str(artifact.get("artifact_id")),
        "preview_type": "table",
        "file_ext": "csv",
        "content_type": artifact.get("content_type"),
        "columns": columns,
        "rows": rows,
        "row_count_previewed": len(rows),
        "truncated": truncated,
        "encoding": encoding,
        "delimiter": delimiter,
        "sheet_names": None,
        "selected_sheet": None,
    }


def xlsx_preview(
    artifact: dict,
    data: bytes,
    *,
    rows_limit: int,
    sheet_name: str | None = None,
    sheet_index: int | None = None,
) -> dict:
    try:
        import openpyxl
    except ImportError as exc:
        raise ArtifactPreviewError(415, "xlsx preview dependency is not installed") from exc

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet_names = list(workbook.sheetnames)
    if not sheet_names:
        selected_name = None
        matrix: list[list[Any]] = []
    else:
        if sheet_name:
            if sheet_name not in sheet_names:
                raise ArtifactPreviewError(400, "sheet_name not found")
            selected_name = sheet_name
        elif sheet_index is not None:
            if sheet_index < 0 or sheet_index >= len(sheet_names):
                raise ArtifactPreviewError(400, "sheet_index out of range")
            selected_name = sheet_names[sheet_index]
        else:
            selected_name = sheet_names[0]
        sheet = workbook[selected_name]
        matrix = []
        for idx, row in enumerate(sheet.iter_rows(values_only=True)):
            matrix.append(list(row))
            if idx > rows_limit:
                break
    columns, rows, truncated = _rows_from_matrix(matrix, rows_limit)
    return {
        "artifact_id": str(artifact.get("artifact_id")),
        "preview_type": "table",
        "file_ext": "xlsx",
        "content_type": artifact.get("content_type"),
        "columns": columns,
        "rows": rows,
        "row_count_previewed": len(rows),
        "truncated": truncated,
        "encoding": None,
        "delimiter": None,
        "sheet_names": sheet_names,
        "selected_sheet": selected_name,
    }


def xls_preview(
    artifact: dict,
    data: bytes,
    *,
    rows_limit: int,
    sheet_name: str | None = None,
    sheet_index: int | None = None,
) -> dict:
    try:
        import xlrd
    except ImportError as exc:
        raise ArtifactPreviewError(415, "xls preview dependency is not installed") from exc

    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    sheet_names = workbook.sheet_names()
    if not sheet_names:
        selected_name = None
        matrix: list[list[Any]] = []
    else:
        if sheet_name:
            if sheet_name not in sheet_names:
                raise ArtifactPreviewError(400, "sheet_name not found")
            selected_name = sheet_name
            sheet = workbook.sheet_by_name(sheet_name)
        else:
            index = sheet_index if sheet_index is not None else 0
            if index < 0 or index >= len(sheet_names):
                raise ArtifactPreviewError(400, "sheet_index out of range")
            selected_name = sheet_names[index]
            sheet = workbook.sheet_by_index(index)
        max_rows = min(sheet.nrows, rows_limit + 2)
        matrix = [sheet.row_values(idx) for idx in range(max_rows)]
    columns, rows, truncated = _rows_from_matrix(matrix, rows_limit)
    return {
        "artifact_id": str(artifact.get("artifact_id")),
        "preview_type": "table",
        "file_ext": "xls",
        "content_type": artifact.get("content_type"),
        "columns": columns,
        "rows": rows,
        "row_count_previewed": len(rows),
        "truncated": truncated,
        "encoding": None,
        "delimiter": None,
        "sheet_names": sheet_names,
        "selected_sheet": selected_name,
    }


def text_preview(artifact: dict, data: bytes, *, text_chars_limit: int) -> dict:
    text, encoding = _decode_text(data)
    truncated = len(text) > text_chars_limit
    return {
        "artifact_id": str(artifact.get("artifact_id")),
        "preview_type": "text",
        "file_ext": infer_preview_ext(artifact),
        "text": text[:text_chars_limit],
        "truncated": truncated,
        "encoding": encoding,
    }


def json_preview(artifact: dict, data: bytes, *, text_chars_limit: int) -> dict:
    text, encoding = _decode_text(data)
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        truncated = len(text) > text_chars_limit
        return {
            "artifact_id": str(artifact.get("artifact_id")),
            "preview_type": "json",
            "file_ext": "json",
            "text": text[:text_chars_limit],
            "truncated": truncated,
            "encoding": encoding,
            "parse_error": str(exc),
        }
    return {
        "artifact_id": str(artifact.get("artifact_id")),
        "preview_type": "json",
        "file_ext": "json",
        "json": parsed,
        "truncated": False,
    }


def pdf_preview(artifact: dict) -> dict:
    artifact_id = str(artifact.get("artifact_id"))
    download_url = f"/artifact-browser/artifacts/{artifact_id}/download"
    return {
        "artifact_id": artifact_id,
        "preview_type": "pdf_inline",
        "file_ext": "pdf",
        "content_type": artifact.get("content_type") or "application/pdf",
        "download_url": download_url,
        "inline_url": f"{download_url}?disposition=inline",
        "note": "Future UI can render this PDF inline from the download URL.",
    }


def build_preview(
    artifact: dict,
    data: bytes | None,
    *,
    rows_limit: int | None = None,
    sheet_name: str | None = None,
    sheet_index: int | None = None,
    text_chars_limit: int | None = None,
) -> dict:
    ext = infer_preview_ext(artifact)
    row_limit = clamp_rows_limit(rows_limit)
    char_limit = clamp_text_chars_limit(text_chars_limit)

    if ext == "pdf":
        return pdf_preview(artifact)
    if data is None:
        raise ArtifactPreviewError(404, "artifact object missing from storage")
    if ext == "csv":
        return csv_preview(artifact, data, rows_limit=row_limit)
    if ext == "xlsx":
        return xlsx_preview(
            artifact,
            data,
            rows_limit=row_limit,
            sheet_name=sheet_name,
            sheet_index=sheet_index,
        )
    if ext == "xls":
        return xls_preview(
            artifact,
            data,
            rows_limit=row_limit,
            sheet_name=sheet_name,
            sheet_index=sheet_index,
        )
    if ext in SUPPORTED_TEXT_EXTS:
        return text_preview(artifact, data, text_chars_limit=char_limit)
    if ext == "json":
        return json_preview(artifact, data, text_chars_limit=char_limit)

    raise ArtifactPreviewError(415, f"Unsupported preview file type: {ext or 'unknown'}")
